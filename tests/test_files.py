import docx, pandas as pd
from pypdf import PdfWriter
from consiz.deterministic import extract_text, file_metadata, folder_metadata


def test_docx_preview(tmp_path):
    p = tmp_path / "a.docx"; d = docx.Document(); d.add_paragraph("Quarterly plan for hiring."); d.add_paragraph("Second para."); d.save(p)
    assert "Quarterly plan" in extract_text(str(p), 500)
    assert file_metadata(str(p))["kind"] == "Word document"


def test_xlsx_preview(tmp_path):
    p = tmp_path / "b.xlsx"; pd.DataFrame({"city": ["A", "B"], "sales": [10, 20]}).to_excel(p, index=False)
    t = extract_text(str(p), 500)
    assert "city" in t and "sales" in t
    assert "1 sheet" in file_metadata(str(p))["extra"]


def test_pdf_metadata(tmp_path):
    p = tmp_path / "c.pdf"; w = PdfWriter(); w.add_blank_page(200, 200); w.add_blank_page(200, 200); w.write(str(p))
    md = file_metadata(str(p))
    assert md["kind"] == "PDF document" and md["extra"] == "2 pages"


def test_binary_has_no_preview(tmp_path):
    p = tmp_path / "x.png"; p.write_bytes(b"\x89PNG\x00\x00")
    assert extract_text(str(p), 100) is None and file_metadata(str(p))["kind"] == "image"


def test_folder_details(tmp_path):
    (tmp_path / "README.md").write_text("# Demo project\nThis folder holds demo scripts.")
    (tmp_path / "src").mkdir(); (tmp_path / "src" / "m.py").write_text("x=1")
    md = folder_metadata(str(tmp_path))
    assert md["direct_files"] == 1 and md["direct_folders"] == 1 and md["files"] == 2
    assert md["created"] and md["subfolders"] == ["src"]
    assert md["previews"][0][0] == "README.md" and "demo scripts" in md["previews"][0][1]
